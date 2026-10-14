#!/usr/bin/env python3
"""Privacy-preserving PowerPoint design inspection (no slide text or media export)."""
import argparse
import hashlib
import json
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

try:
    from pptx import Presentation
except ImportError as exc:
    raise SystemExit('Install optional inspection dependency: pip install "python-pptx>=1,<2"') from exc

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def inspect_deck(source: Path) -> dict:
    prs = Presentation(str(source))
    sha = hashlib.sha256()
    with source.open('rb') as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b''):
            sha.update(chunk)
    shapes = Counter()
    fonts = Counter()
    image_count = 0
    hidden = []
    for slide_num, slide in enumerate(prs.slides, 1):
        if slide._element.get('show') == '0':
            hidden.append(slide_num)
        for shape in slide.shapes:
            shapes[str(shape.shape_type).split(' (')[0]] += 1
            if str(shape.shape_type).startswith('PICTURE'):
                image_count += 1
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.text.strip() and run.font.name:
                            fonts[run.font.name] += len(run.text)
    theme = {}
    with ZipFile(source) as zf:
        if 'ppt/theme/theme1.xml' in zf.namelist():
            root = ET.fromstring(zf.read('ppt/theme/theme1.xml'))
            schemes = root.find('.//' + A + 'clrScheme')
            if schemes is not None:
                for entry in schemes:
                    if len(entry):
                        value = entry[0].get('val') or entry[0].get('lastClr')
                        if value:
                            theme[entry.tag.split('}')[-1]] = '#' + value if len(value)==6 else value
            tf = root.find('.//' + A + 'fontScheme')
            if tf is not None:
                for role in ('majorFont', 'minorFont'):
                    el = tf.find(A+role)
                    if el is not None:
                        latin = el.find(A+'latin')
                        if latin is not None:
                            theme[role + 'Latin'] = latin.get('typeface')
    return {
        'sourceFileName': source.name,
        'sourceSHA256': sha.hexdigest(),
        'privacy': 'No slide text, notes, images or binary data exported',
        'slideCount': len(prs.slides),
        'hiddenSlideIndices': hidden,
        'visibleSlideCount': len(prs.slides)-len(hidden),
        'slideDimensionsInches': {
            'width': round(prs.slide_width / 914400, 3),
            'height': round(prs.slide_height / 914400, 3),
        },
        'slideMasters': len(prs.slide_masters),
        'layouts': [layout.name for layout in prs.slide_layouts],
        'shapeCounts': dict(sorted(shapes.items())),
        'imageOccurrences': image_count,
        'observedFonts': sorted(fonts),
        'theme': theme,
        'reviewStatus': 'Extracted technical metadata only; brand and visual approval pending'
    }


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('pptx', type=Path)
    ap.add_argument('--out', type=Path, required=True)
    args = ap.parse_args()
    if args.pptx.suffix.lower() != '.pptx':
        ap.error('Expected an editable .pptx input')
    data = inspect_deck(args.pptx)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(data, indent=2) + '\n', encoding='utf8')
    print(f"Inspected {data['slideCount']} slides, {len(data['layouts'])} layouts. Privacy-safe metadata: {args.out}")

if __name__ == '__main__':
    main()
